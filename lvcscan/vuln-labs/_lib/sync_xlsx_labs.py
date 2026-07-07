#!/usr/bin/env python3
"""Sync Vulnerable Lab column in docs/Laravel_Vulnerabilities.xlsx from on-disk run.sh labs."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "docs" / "Laravel_Vulnerabilities.xlsx"

# Explicit multi-lab rows (CVE -> list of relative paths)
MULTI = {
    "CVE-2024-21546": [
        "vuln-labs/third-party/unisharp-laravel-filemanager/2.8.1/cve-2024-21546_unauth-filemanager-upload-rce_35002-36002",
        "vuln-labs/third-party/unisharp-laravel-filemanager/2.8.1/cve-2024-21546_auth-filemanager-upload-rce_35001-36001",
        "vuln-labs/apps/badaso/2.9.10/cve-2024-21546_auth-filemanager-upload-rce_41013-42013",
    ],
}

# Map CVE id (normalized) -> lab dir relative path (from run.sh under vuln-labs/, skip medium/)
def discover_labs() -> dict[str, list[str]]:
    by_cve: dict[str, list[str]] = {}
    lab_roots = (ROOT / "vuln-labs",)
    for lab_root in lab_roots:
        if not lab_root.is_dir():
            continue
        for run in lab_root.rglob("run.sh"):
            if "medium" in run.parts:
                continue
            rel_dir = run.parent.relative_to(ROOT).as_posix()
            m = re.search(r"(cve-\d{4}-\d+)", rel_dir, re.I)
            if not m:
                continue
            slug = m.group(1).upper().replace("CVE-", "CVE-")
            # normalize CVE-YYYY-NNNNN
            parts = slug.split("-")
            cve = f"{parts[0]}-{parts[1]}-{parts[2]}"
            by_cve.setdefault(cve, []).append(rel_dir)
    for k in by_cve:
        by_cve[k] = sorted(set(by_cve[k]))
    by_cve.update(MULTI)
    return by_cve


def main() -> None:
    labs = discover_labs()
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = headers.index("Vulnerable Lab") + 1
    updated = 0
    for r in range(2, ws.max_row + 1):
        cve = ws.cell(r, 1).value
        if not cve:
            continue
        cve = str(cve).strip()
        if cve in labs:
            val = "\n".join(labs[cve])
            old = ws.cell(r, col).value
            if old != val:
                ws.cell(r, col).value = val
                updated += 1
                print(f"update {cve}")
        else:
            print(f"MISSING lab for {cve}")
    wb.save(XLSX)
    print(f"\nSaved {XLSX.name}; {updated} rows updated.")


if __name__ == "__main__":
    main()
