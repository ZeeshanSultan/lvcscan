#!/usr/bin/env python3
"""Sync Exploit Pre-conditions column in docs/Laravel_Vulnerabilities.xlsx from CVE_METADATA."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
XLSX = ROOT / "docs" / "Laravel_Vulnerabilities.xlsx"


def main() -> None:
    from modules.cves.metadata import CVE_METADATA

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = headers.index("Exploit Pre-conditions") + 1
    updated = 0
    for r in range(2, ws.max_row + 1):
        cve = ws.cell(r, 1).value
        if not cve:
            continue
        cve = str(cve).strip()
        meta = CVE_METADATA.get(cve)
        if not meta:
            print(f"SKIP {cve}: not in CVE_METADATA")
            continue
        val = "\n".join(meta.preconditions or []).strip()
        old = ws.cell(r, col).value
        if (old or "").strip() != val:
            ws.cell(r, col).value = val
            updated += 1
            print(f"update {cve}")
    wb.save(XLSX)
    print(f"\nSaved {XLSX.name}; {updated} precondition rows updated.")


if __name__ == "__main__":
    main()
